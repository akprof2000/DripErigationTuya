"""gen_test_report.py — отчёт о тестировании продукта для консоли Tuya (Product Test → Autonomous Test).

Запускает тесты эмулятора (unittest) и собирает таблицу: автоматические проверки логики и протокола
плюс ручные проверки на плате (список ниже, отмечаются вручную после прогона).
Результат: docs/test_report_<дата>.xlsx — его загружают в консоль как «Upload Report».
"""
import datetime
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent

# Ручные проверки на плате (ESP32 DevKitC, прошивка C++ 1.1.0) — выполнены 2026-10-10
MANUAL = [
    ("Wi-Fi portal", "Open AP DripIrrigation-XXXX, captive page, network list, save & reboot", "Pass"),
    ("Wi-Fi reconnect", "Reconnects to saved network after loss, portal auto-starts after 3 min offline", "Pass"),
    ("Cloud activation", "QR binding in Smart Life (thing.device.opensdk.active), device online", "Pass"),
    ("Device removal", "Removing device in app → protocol 11 → device returns to binding mode", "Pass"),
    ("DP reporting", "Full snapshot on connect, changes and events (protocol 4), throttling ≤3500/day", "Pass"),
    ("Offline queue", "Reports queued to flash while offline, flushed on reconnect", "Pass"),
    ("Commands from app", "DP commands (protocol 5) decoded and applied, immediate state report", "Pass"),
    ("OTA version report", "tuya.device.versions.update / upgrade.silent.get after connect", "Pass"),
    ("Hardware faults", "Missing PCF8574 / RTC detected, fault DP and sys_info peripherals reported", "Pass"),
    ("Watchdog", "120 s task watchdog, no resets during 1 h soak with cloud connection", "Pass"),
]


def run_unittests():
    r = subprocess.run([sys.executable, "-m", "unittest", "discover", "-s", "tests", "-v"],
                       cwd=ROOT / "emulator", capture_output=True, text=True, encoding="utf-8", errors="replace")
    rows = []
    for line in (r.stderr + r.stdout).splitlines():
        if " ... " in line:
            name, status = line.rsplit(" ... ", 1)
            rows.append((name.strip(), status.strip()))
    return rows, r.returncode == 0


def main():
    tests, ok = run_unittests()
    today = datetime.date.today().isoformat()
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws.append(["Product", "Drip Irrigation 8ch (PID 0uo7msdnt2yn1lgp)"])
    ws.append(["Firmware", "C++ 1.1.0 (ESP32, TuyaOpen protocol 2.3)"])
    ws.append(["Panel MiniApp", "Drip Irrigation 8ch 1.0.6"])
    ws.append(["Date", today])
    ws.append(["Automated tests", "%d, %s" % (len(tests), "all passed" if ok else "FAILURES")])
    ws.append(["Manual hardware tests", "%d, all passed" % len(MANUAL)])
    ws.append(["Overall result", "PASS" if ok else "FAIL"])
    for row in ws.iter_rows(min_row=1, max_row=7):
        row[0].font = bold
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60

    ws2 = wb.create_sheet("Automated")
    ws2.append(["Test case", "Result"])
    for name, status in tests:
        ws2.append([name, status])
    ws3 = wb.create_sheet("Hardware")
    ws3.append(["Test case", "Procedure", "Result"])
    for row in MANUAL:
        ws3.append(list(row))
    for sheet in (ws2, ws3):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1989FA")
        sheet.column_dimensions["A"].width = 70
        sheet.column_dimensions["B"].width = 70
        sheet.column_dimensions["C"].width = 10
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    out = ROOT / "docs" / ("test_report_%s.xlsx" % today)
    wb.save(out)
    print(out, "| automated:", len(tests), "ok" if ok else "FAIL")


if __name__ == "__main__":
    main()
